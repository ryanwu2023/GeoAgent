"""Meeting calendar filtering and safe XLSX export helpers."""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Any
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "事件", "相关人物（职务）", "国家/地区", "会晤地点", "开始日期", "结束日期",
    "日期精度", "当前状态", "主题", "事件类型", "是否非会晤政治日程", "是否存在口径冲突",
    "来源数量", "来源名称", "原始来源链接", "首次发现时间", "最后更新时间", "备注",
]

TOPIC_LABELS = {"all": "全部外交日程", "usiran": "美伊相关日程", "ukraine": "俄乌相关日程"}
PRECISION_LABELS = {
    "day": "明确日期", "range": "日期范围", "week": "周", "month": "月份",
    "quarter": "季度", "year": "年份", "unknown": "待确认",
}
FAMILY_LABELS = {
    "meeting": "会晤", "summit": "峰会", "visit": "访问", "election": "选举",
    "conference": "会议", "political_schedule": "政治日程",
}


def _parse_anchor(value: Any) -> date:
    try:
        return date.fromisoformat(str(value or "")[:10])
    except ValueError:
        return datetime.now(timezone.utc).date()


def filter_meetings(
    snapshot: dict[str, Any], *, topic: str = "all", track: str = "", status: str = "",
    person: str = "", start: str = "", end: str = "", include_inactive: bool = False,
) -> tuple[list[dict[str, Any]], str, str]:
    """Apply the same filters used by the meeting page and export endpoint."""
    anchor = _parse_anchor(snapshot.get("created_at"))
    range_start = start or anchor.isoformat()
    range_end = end or (anchor + timedelta(days=120)).isoformat()
    requested_statuses = {value.strip() for value in status.split(",") if value.strip()}
    explicit_range = bool(start or end)
    rows: list[dict[str, Any]] = []
    for item in snapshot.get("meetings", []):
        if topic != "all" and topic not in item.get("topics", []):
            continue
        if track and item.get("track") != track:
            continue
        if requested_statuses and item.get("status") not in requested_statuses:
            continue
        if person and not any(
            person.casefold() in str(value).casefold()
            for entry in item.get("persons", [])
            for value in (entry.get("id", ""), entry.get("name", ""))
        ):
            continue
        if not include_inactive and item.get("inactive"):
            continue
        event_start = item.get("date_start") or ""
        event_end = item.get("date_end") or event_start
        if event_start and (event_end < range_start or event_start > range_end):
            continue
        if not event_start and explicit_range:
            continue
        rows.append(item)
    rows.sort(key=lambda item: (item.get("date_start") or "9999-12-31", item.get("id") or ""))
    return rows, range_start, range_end


def _safe_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    if len(text) > 32_000:
        text = text[:31_990] + "…（已截断）"
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def _valid_url(value: Any) -> bool:
    try:
        parsed = urlsplit(str(value or "").strip())
        return parsed.scheme.lower() in {"http", "https"} and bool(parsed.netloc)
    except ValueError:
        return False


def _excel_date(value: Any) -> date | str:
    try:
        return date.fromisoformat(str(value or "")[:10])
    except ValueError:
        return _safe_cell(value)


def _topic_text(item: dict[str, Any]) -> str:
    labels = [TOPIC_LABELS.get(topic, str(topic)) for topic in item.get("topics", [])]
    return "、".join(labels) or "未分类"


def _people_text(item: dict[str, Any]) -> str:
    people = []
    for person in item.get("persons", []):
        name = _safe_cell(person.get("name") or person.get("id") or "")
        role = _safe_cell(person.get("role") or "")
        if name:
            people.append(f"{name}（{role}）" if role else name)
    return "、".join(people) or "参与人待披露"


def _countries_text(item: dict[str, Any]) -> str:
    values = []
    for person in item.get("persons", []):
        country = _safe_cell(person.get("country") or "")
        if country and country not in values:
            values.append(country)
    return "、".join(values) or "待补充"


def _remarks(item: dict[str, Any]) -> str:
    notes = [item.get("baseline_note"), item.get("date_raw")]
    if item.get("conflict"):
        notes.append("不同来源存在口径冲突，请逐条核对")
    if item.get("non_meeting"):
        notes.append("政治日程，非会晤")
    return "；".join(_safe_cell(note) for note in notes if note)


def render_meeting_workbook(rows: list[dict[str, Any]], metadata: dict[str, Any]) -> bytes:
    """Render a formula-free, styled XLSX workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "会晤台账"
    sheet.append(HEADERS)

    for item in rows:
        sources = item.get("sources", [])
        source_names = [source.get("name") or source.get("outlet") or "公开来源" for source in sources]
        urls = [_safe_cell(source.get("url")) for source in sources if _valid_url(source.get("url"))]
        row_values: list[Any] = [
            _safe_cell(item.get("title") or "未命名事项"), _people_text(item), _countries_text(item),
            _safe_cell(item.get("location") or "地点未定"), _excel_date(item.get("date_start")),
            _excel_date(item.get("date_end")), PRECISION_LABELS.get(item.get("precision"), _safe_cell(item.get("precision") or "待确认")),
            _safe_cell(item.get("status_label") or item.get("status") or "待确认"), _topic_text(item),
            FAMILY_LABELS.get(item.get("family"), _safe_cell(item.get("track_label") or item.get("family") or "会晤")),
            "是" if item.get("non_meeting") else "否", "是" if item.get("conflict") else "否",
            len(sources), "、".join(_safe_cell(name) for name in source_names), "\n".join(urls),
            _safe_cell(item.get("first_seen")), _safe_cell(item.get("last_seen")), _remarks(item),
        ]
        sheet.append(row_values)
        link_cell = sheet.cell(sheet.max_row, 15)
        if urls:
            link_cell.hyperlink = urls[0]
            link_cell.style = "Hyperlink"

    header_fill = PatternFill("solid", fgColor="17365D")
    alternate_fill = PatternFill("solid", fgColor="EAF2F8")
    white_font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
    body_font = Font(name="Microsoft YaHei", color="1F2937", size=10)
    thin = Side(style="thin", color="D5DEE7")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row_number in range(2, sheet.max_row + 1):
        for cell in sheet[row_number]:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if row_number % 2 == 0:
                cell.fill = alternate_fill
        for column in (5, 6):
            sheet.cell(row_number, column).number_format = "yyyy-mm-dd"
    widths = [30, 28, 16, 18, 13, 13, 13, 14, 20, 18, 17, 17, 11, 24, 42, 22, 22, 38]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{max(sheet.max_row, 1)}"
    sheet.sheet_view.showGridLines = False

    info = workbook.create_sheet("导出说明")
    info_rows = [
        ("导出说明", "值"),
        ("导出时间", metadata.get("exported_at", "")),
        ("数据截至", metadata.get("as_of", "")),
        ("主题", metadata.get("topic", "全部外交日程")),
        ("状态", metadata.get("status") or "全部状态"),
        ("人物", metadata.get("person") or "全部人物"),
        ("日期范围", f"{metadata.get('start', '')} 至 {metadata.get('end', '')}"),
        ("包含非活动事项", "是" if metadata.get("include_inactive") else "否"),
        ("导出记录数", len(rows)),
        ("数据边界", metadata.get("basis") or "冻结快照中的结构化会晤日历；计划、媒体报道和人工基线不等于会晤必然举行。"),
    ]
    for row in info_rows:
        info.append([_safe_cell(row[0]), row[1] if isinstance(row[1], int) else _safe_cell(row[1])])
    for cell in info[1]:
        cell.fill = header_fill
        cell.font = white_font
    for row in info.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 86
    info.freeze_panes = "A2"
    info.sheet_view.showGridLines = False

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def meeting_export_filename(exported_at: datetime | None = None) -> str:
    current = exported_at or datetime.now(timezone.utc)
    beijing = current.astimezone(timezone(timedelta(hours=8)))
    return f"GeoAgent会晤日历_{beijing.date().isoformat()}.xlsx"
