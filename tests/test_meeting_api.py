from io import BytesIO

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from backend.app import app
from backend.store import save_snapshot


client = TestClient(app)


def meeting(event_id, *, status="confirmed", topics=None, date_start="2026-10-27", location="莫斯科"):
    return {
        "id": event_id, "evidence_id": "meeting-" + event_id, "title": "领导人会晤",
        "track": "us-ir", "topics": topics or ["usiran"], "date_start": date_start,
        "date_end": "", "precision": "day", "location": location, "points": [],
        "persons": [{"id": "putin", "name": "普京", "role": "俄罗斯总统", "country": "俄罗斯"}],
        "status": status, "status_label": "已确认", "inactive": status in {"canceled", "lapsed"},
        "claims": [], "sources": [], "updates": [], "conflict": False,
    }


def snapshot_payload(meetings):
    return {"schema_version": "geo.v6/1", "records": [], "observations": [], "sources": [],
            "meetings": meetings, "manifest": [], "errors": [], "projects": [], "cleaning": {}}


def test_meeting_api_filters_topic_person_status_and_inactive(tmp_path, monkeypatch):
    monkeypatch.setenv("GEO_DATA_DIR", str(tmp_path))
    sid, _ = save_snapshot(snapshot_payload([
        meeting("EV-1"), meeting("EV-2", status="canceled"),
        meeting("EV-3", topics=["ukraine"], date_start="2026-11-02"),
    ]))
    response = client.get("/api/meetings", params={"snapshot_id": sid, "topic": "usiran", "person": "普京"})
    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 1
    assert payload["records"][0]["id"] == "EV-1"
    assert payload["status_counts"] == {"confirmed": 1}
    with_inactive = client.get("/api/meetings", params={"snapshot_id": sid, "topic": "usiran", "include_inactive": True}).json()
    assert {item["id"] for item in with_inactive["records"]} == {"EV-1", "EV-2"}


def test_meeting_detail_is_snapshot_scoped(tmp_path, monkeypatch):
    monkeypatch.setenv("GEO_DATA_DIR", str(tmp_path))
    sid, _ = save_snapshot(snapshot_payload([meeting("EV-1")]))
    other, _ = save_snapshot(snapshot_payload([meeting("EV-2")]))
    assert client.get("/api/meetings/EV-1", params={"snapshot_id": sid}).status_code == 200
    assert client.get("/api/meetings/EV-1", params={"snapshot_id": other}).status_code == 404
    assert client.get("/api/meetings/EV-X", params={"snapshot_id": sid}).status_code == 404


def test_meeting_api_honors_explicit_date_range(tmp_path, monkeypatch):
    monkeypatch.setenv("GEO_DATA_DIR", str(tmp_path))
    sid, _ = save_snapshot(snapshot_payload([meeting("EV-1"), meeting("EV-2", date_start="2026-12-01")]))
    payload = client.get("/api/meetings", params={
        "snapshot_id": sid, "start": "2026-10-25", "end": "2026-11-01",
    }).json()
    assert [item["id"] for item in payload["records"]] == ["EV-1"]


def test_meeting_export_uses_current_filters_and_download_headers(tmp_path, monkeypatch):
    monkeypatch.setenv("GEO_DATA_DIR", str(tmp_path))
    sid, _ = save_snapshot(snapshot_payload([
        meeting("EV-1"), meeting("EV-2", topics=["ukraine"], date_start="2026-10-28"),
    ]))
    response = client.get("/api/meetings/export.xlsx", params={
        "snapshot_id": sid, "topic": "usiran", "person": "putin",
        "start": "2026-10-27", "end": "2026-10-27",
    })
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook["会晤台账"].max_row == 2
    assert workbook["会晤台账"]["A2"].value == "领导人会晤"
    assert workbook["导出说明"]["B9"].value == 1


def test_meeting_export_returns_conflict_when_filter_is_empty(tmp_path, monkeypatch):
    monkeypatch.setenv("GEO_DATA_DIR", str(tmp_path))
    sid, _ = save_snapshot(snapshot_payload([meeting("EV-1")]))
    response = client.get("/api/meetings/export.xlsx", params={
        "snapshot_id": sid, "start": "2027-01-01", "end": "2027-01-02",
    })
    assert response.status_code == 409
    assert response.json()["detail"] == "当前筛选范围没有可导出的会晤日程"
